"""RTO, Make/Model and Pincode master dashboards: search endpoints, create modals, exports and edit pages."""
from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from django.core.cache import cache
import json
from openpyxl import Workbook
from ..models import RTOMaster, MakeModelMaster, PincodeMaster, AuditLog
from ..forms import MakeModelForm, PincodeForm, RTOForm
from ..services.masters import RTO_MAKE_CHOICES_CACHE_KEY, _sync_pincode_zones_from_health_data


# -------------------------
# RTO DASHBOARD
# -------------------------
RTO_DASHBOARD_DEFAULT_LIMIT = 20


def _filtered_rto_qs(request):
    qs = RTOMaster.objects.all().order_by("rto_name")

    rto_names = request.GET.getlist("rto_name")
    cluster_q = (request.GET.get("cluster_q") or "").strip()

    if rto_names and "" not in rto_names:
        qs = qs.filter(rto_name__in=rto_names)
    if cluster_q:
        qs = qs.filter(rto_cluster__icontains=cluster_q)

    return qs


def rto_dashboard(request):
    """
    Renders the page shell only. The table itself is populated by JS via
    rto_dashboard_search (default: the 20 most recently added/updated rows),
    so the initial page load no longer pulls every RTOMaster row regardless
    of whether the user ever searches.
    """
    rto_name_list = RTOMaster.objects.values_list("rto_name", flat=True).distinct().order_by("rto_name")

    return render(request, "rto_dashboard.html", {
        "total": RTOMaster.objects.count(),
        "rto_name_list": rto_name_list,
        "default_limit": RTO_DASHBOARD_DEFAULT_LIMIT,
        "is_admin": True
    })


def rto_dashboard_search(request):
    """
    JSON backing endpoint for the dashboard table. With a `limit` param it
    returns just that many rows, most recently added/updated first (the
    page's own default view, and "Reset"); without one it returns every row
    matching the filters, unbounded -- only reached when the user actually
    clicks Apply.
    """
    qs = _filtered_rto_qs(request)

    limit_raw = (request.GET.get("limit") or "").strip()
    if limit_raw:
        try:
            limit = max(1, int(limit_raw))
        except ValueError:
            limit = RTO_DASHBOARD_DEFAULT_LIMIT
        qs = qs.order_by("-updated_at", "-id")[:limit]

    results = [
        {"id": r.id, "rto_name": r.rto_name, "rto_cluster": r.rto_cluster or ""}
        for r in qs
    ]
    return JsonResponse({"results": results, "total": len(results)})


def create_rto(request):
    """
    "Create New Group" endpoint for the dashboard's modal -- adds a single
    new RTOMaster row without leaving the page. Same iexact pre-check as
    api_upload_chunk's rto_master branch, so a case variant of an existing
    name gets a friendly message instead of an IntegrityError.
    """
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid request method."}, status=405)

    try:
        data = json.loads(request.body)
    except (json.JSONDecodeError, TypeError):
        return JsonResponse({"success": False, "message": "Invalid request body."}, status=400)

    name = (data.get("rto_name") or "").strip()
    cluster = (data.get("rto_cluster") or "").strip()

    if not name:
        return JsonResponse({"success": False, "message": "Group name is required."}, status=400)
    if len(name) > 100:
        return JsonResponse({"success": False, "message": "Group name is longer than 100 characters."}, status=400)
    if RTOMaster.objects.filter(rto_name__iexact=name).exists():
        return JsonResponse({"success": False, "message": f'"{name}" already exists.'}, status=400)

    obj = RTOMaster.objects.create(rto_name=name, rto_cluster=cluster or None)

    AuditLog.objects.create(
        user=request.user,
        action="RTO MASTER CREATE",
        details=(
            f"Created new RTO Master group '{name}' (id {obj.id}) from the RTO Dashboard."
            + (f" Cluster: {cluster}" if cluster else "")
        ),
    )
    # Mirrors api_upload_chunk / create_make_model: every write to this master
    # table busts the key get_rto_and_make_choices caches the rate forms'
    # dropdowns under.
    cache.delete(RTO_MAKE_CHOICES_CACHE_KEY)

    return JsonResponse({
        "success": True,
        "id": obj.id,
        "rto_name": obj.rto_name,
        "rto_cluster": obj.rto_cluster or "",
    })


def export_rto_xlsx(request):
    qs = _filtered_rto_qs(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "RTO Master"
    ws.append(["ID", "RTO NAME", "RTO CLUSTER"])

    for r in qs:
        ws.append([r.id, r.rto_name, r.rto_cluster or ""])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="rto_master.xlsx"'
    wb.save(response)
    return response


# -------------------------
# MAKE MODEL DASHBOARD
# -------------------------
MAKE_MODEL_DASHBOARD_DEFAULT_LIMIT = 20


def _filtered_make_model_qs(request):
    qs = MakeModelMaster.objects.all().order_by("make_model_name")

    make_model_names = request.GET.getlist("make_model_name")
    cluster_q = (request.GET.get("cluster_q") or "").strip()

    if make_model_names and "" not in make_model_names:
        qs = qs.filter(make_model_name__in=make_model_names)
    if cluster_q:
        qs = qs.filter(make_model_cluster__icontains=cluster_q)

    return qs


def make_model_dashboard(request):
    """
    Renders the page shell only. The table itself is populated by JS via
    make_model_dashboard_search (default: the 20 most recently added/updated
    rows), so the initial page load no longer pulls every MakeModelMaster
    row regardless of whether the user ever searches.
    """
    make_model_name_list = MakeModelMaster.objects.values_list(
        "make_model_name", flat=True
    ).distinct().order_by("make_model_name")

    return render(request, "make_model_dashboard.html", {
        "total": MakeModelMaster.objects.count(),
        "make_model_name_list": make_model_name_list,
        "default_limit": MAKE_MODEL_DASHBOARD_DEFAULT_LIMIT,
        "is_admin": True
    })


def make_model_dashboard_search(request):
    """
    JSON backing endpoint for the dashboard table. With a `limit` param it
    returns just that many rows, most recently added/updated first (the
    page's own default view, and "Reset"); without one it returns every row
    matching the filters, unbounded -- only reached when the user actually
    clicks Apply.
    """
    qs = _filtered_make_model_qs(request)

    limit_raw = (request.GET.get("limit") or "").strip()
    if limit_raw:
        try:
            limit = max(1, int(limit_raw))
        except ValueError:
            limit = MAKE_MODEL_DASHBOARD_DEFAULT_LIMIT
        qs = qs.order_by("-updated_at", "-id")[:limit]

    results = [
        {
            "id": r.id,
            "make_model_name": r.make_model_name,
            "make_model_cluster": r.make_model_cluster or "",
        }
        for r in qs
    ]
    return JsonResponse({"results": results, "total": len(results)})


def create_make_model(request):
    """
    "Create New Group" endpoint for the dashboard's modal -- adds a single
    new MakeModelMaster row without leaving the page. Same iexact pre-check
    as add_missing_make_model_to_master's "new" target, so a case variant of
    an existing name gets a friendly message instead of an IntegrityError.
    """
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid request method."}, status=405)

    try:
        data = json.loads(request.body)
    except (json.JSONDecodeError, TypeError):
        return JsonResponse({"success": False, "message": "Invalid request body."}, status=400)

    name = (data.get("make_model_name") or "").strip()
    cluster = (data.get("make_model_cluster") or "").strip()

    if not name:
        return JsonResponse({"success": False, "message": "Group name is required."}, status=400)
    if len(name) > 150:
        return JsonResponse({"success": False, "message": "Group name is longer than 150 characters."}, status=400)
    if MakeModelMaster.objects.filter(make_model_name__iexact=name).exists():
        return JsonResponse({"success": False, "message": f'"{name}" already exists.'}, status=400)

    obj = MakeModelMaster.objects.create(make_model_name=name, make_model_cluster=cluster or None)

    AuditLog.objects.create(
        user=request.user,
        action="MAKE MODEL MASTER CREATE",
        details=(
            f"Created new Make/Model Master group '{name}' (id {obj.id}) from the Make Model Dashboard."
            + (f" Cluster: {cluster}" if cluster else "")
        ),
    )
    # Mirrors api_upload_chunk / add_missing_make_model_to_master: every write
    # to this master table busts the key get_rto_and_make_choices caches the
    # rate forms' dropdowns under.
    cache.delete(RTO_MAKE_CHOICES_CACHE_KEY)

    return JsonResponse({
        "success": True,
        "id": obj.id,
        "make_model_name": obj.make_model_name,
        "make_model_cluster": obj.make_model_cluster or "",
    })


def export_make_model_xlsx(request):
    qs = _filtered_make_model_qs(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Make Model Master"
    ws.append(["ID", "MAKE MODEL NAME", "MAKE MODEL CLUSTER"])

    for r in qs:
        ws.append([r.id, r.make_model_name, r.make_model_cluster or ""])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="make_model_master.xlsx"'
    wb.save(response)
    return response


PINCODE_DASHBOARD_DEFAULT_LIMIT = 20


def _filtered_pincode_qs(request):
    qs = PincodeMaster.objects.all().order_by("pincode_zone")

    zone_names = request.GET.getlist("pincode_zone")
    cluster_q = (request.GET.get("cluster_q") or "").strip()

    if zone_names and "" not in zone_names:
        qs = qs.filter(pincode_zone__in=zone_names)
    if cluster_q:
        qs = qs.filter(pincode_cluster__icontains=cluster_q)

    return qs


def pincode_dashboard(request):
    """
    Renders the page shell only. The table itself is populated by JS via
    pincode_dashboard_search (default: the 20 most recently added/updated
    rows), so the initial page load no longer pulls every PincodeMaster row
    regardless of whether the user ever searches.
    """
    _sync_pincode_zones_from_health_data()

    zone_name_list = PincodeMaster.objects.values_list("pincode_zone", flat=True).distinct().order_by("pincode_zone")

    return render(request, "pincode_dashboard.html", {
        "total": PincodeMaster.objects.count(),
        "zone_name_list": zone_name_list,
        "default_limit": PINCODE_DASHBOARD_DEFAULT_LIMIT,
        "is_admin": True
    })


def pincode_dashboard_search(request):
    """
    JSON backing endpoint for the dashboard table. With a `limit` param it
    returns just that many rows, most recently added/updated first (the
    page's own default view, and "Reset"); without one it returns every row
    matching the filters, unbounded -- only reached when the user actually
    clicks Apply. Re-syncs zones first, same as every other entry point into
    this table, so a zone that only just appeared in HealthRateMaster is
    searchable immediately rather than after the next full page load.
    """
    _sync_pincode_zones_from_health_data()

    qs = _filtered_pincode_qs(request)

    limit_raw = (request.GET.get("limit") or "").strip()
    if limit_raw:
        try:
            limit = max(1, int(limit_raw))
        except ValueError:
            limit = PINCODE_DASHBOARD_DEFAULT_LIMIT
        qs = qs.order_by("-updated_at", "-id")[:limit]

    results = [
        {"id": r.id, "pincode_zone": r.pincode_zone, "pincode_cluster": r.pincode_cluster or ""}
        for r in qs
    ]
    return JsonResponse({"results": results, "total": len(results)})


def create_pincode(request):
    """
    "Create New Group" endpoint for the dashboard's modal -- adds a single
    new PincodeMaster row without leaving the page. Same iexact pre-check as
    the RTO/Make-Model create endpoints, so a case variant of an existing
    zone gets a friendly message instead of an IntegrityError.
    """
    if request.method != "POST":
        return JsonResponse({"success": False, "message": "Invalid request method."}, status=405)

    try:
        data = json.loads(request.body)
    except (json.JSONDecodeError, TypeError):
        return JsonResponse({"success": False, "message": "Invalid request body."}, status=400)

    zone = (data.get("pincode_zone") or "").strip()
    cluster = (data.get("pincode_cluster") or "").strip()

    if not zone:
        return JsonResponse({"success": False, "message": "Zone name is required."}, status=400)
    if len(zone) > 100:
        return JsonResponse({"success": False, "message": "Zone name is longer than 100 characters."}, status=400)
    if PincodeMaster.objects.filter(pincode_zone__iexact=zone).exists():
        return JsonResponse({"success": False, "message": f'"{zone}" already exists.'}, status=400)

    obj = PincodeMaster.objects.create(pincode_zone=zone, pincode_cluster=cluster or None)

    AuditLog.objects.create(
        user=request.user,
        action="PINCODE MASTER CREATE",
        details=(
            f"Created new Pincode Master zone '{zone}' (id {obj.id}) from the Pincode Dashboard."
            + (f" Cluster: {cluster}" if cluster else "")
        ),
    )

    return JsonResponse({
        "success": True,
        "id": obj.id,
        "pincode_zone": obj.pincode_zone,
        "pincode_cluster": obj.pincode_cluster or "",
    })


def export_pincode_xlsx(request):
    _sync_pincode_zones_from_health_data()

    qs = _filtered_pincode_qs(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pincode Master"
    # Header row matches PincodeMaster field names — lets this export be
    # edited and re-uploaded via Import Data's "Pincode Master (Clusters CSV)"
    # option without renaming any columns.
    ws.append(["id", "pincode_zone", "pincode_cluster"])

    for r in qs:
        ws.append([r.id, r.pincode_zone, r.pincode_cluster or ""])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="pincode_master.xlsx"'
    wb.save(response)
    return response


# -------------------------
# Edit Master Tables
# -------------------------
def edit_rto(request, pk):
    obj = RTOMaster.objects.get(id=pk)
    if request.method == "POST":
        form = RTOForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            return redirect("rto_dashboard")
    else:
        form = RTOForm(instance=obj)

    return render(request, "edit_master.html", {
        "form": form,
        "title": "Edit RTO Record",
        "back_url": "rto_dashboard"
    })


def edit_make_model(request, pk):
    obj = MakeModelMaster.objects.get(id=pk)
    if request.method == "POST":
        form = MakeModelForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            return redirect("make_model_dashboard")
    else:
        form = MakeModelForm(instance=obj)

    return render(request, "edit_master.html", {
        "form": form,
        "title": "Edit Make/Model Record",
        "back_url": "make_model_dashboard"
    })


def edit_pincode(request, pk):
    obj = PincodeMaster.objects.get(id=pk)
    if request.method == "POST":
        form = PincodeForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            return redirect("pincode_dashboard")
    else:
        form = PincodeForm(instance=obj)

    return render(request, "edit_master.html", {
        "form": form,
        "title": "Edit Pincode Zone Record",
        "back_url": "pincode_dashboard"
    })
