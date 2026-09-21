import json
from unittest import mock

from django.contrib.auth.models import Group, User
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.urls.converters import get_converters
from rest_framework_api_key.models import APIKey
from rest_framework_api_key.permissions import HasAPIKey

from insurance import sso
from insurance import urls as insurance_urls

# Routes in insurance/urls.py that are deliberately reachable while logged
# out (the password-reset flow, and the catch-all which just bounces
# everyone to login regardless of auth state). Every other name in that
# file is expected to refuse an anonymous request — this test exists so a
# newly added path() that forgets its login_required/page_access_required/
# staff_required wrapper fails loudly in CI instead of silently shipping as
# a public page.
#
# Server-to-server API routes are deliberately NOT listed here. They are
# gated by their own HasAPIKey permission rather than a session, which the
# test detects directly (see _api_key_gated) — so they are *verified* to
# reject a keyless request rather than skipped, and a new one needs no entry
# here. Only genuinely public pages belong in this set.
PUBLIC_URL_NAMES = {
    "password_reset",
    "password_reset_done",
    "password_reset_confirm",
    "password_reset_complete",
    "catch_all",
    # The landing page a not-yet-logged-in browser is redirected to by the
    # inbound partner-portal SSO handoff (insurance/sso.py) — meant to be
    # reachable while logged out. Its sibling, sso_issue_ticket, is covered
    # by the HasAPIKey path above instead.
    "sso_consume",
}

# Registered converter name (e.g. "int", "str", "uuid") -> a dummy value
# that satisfies it, so every named URL can be reversed without any real
# objects existing in the test DB.
_DUMMY_VALUE_BY_CONVERTER = {
    "int": "1",
    "str": "dummy",
    "slug": "dummy-slug",
    "uuid": "00000000-0000-0000-0000-000000000000",
    "path": "dummy/path",
}

_CONVERTER_NAME_BY_CLASS = {type(inst): name for name, inst in get_converters().items()}


def _dummy_kwargs(url_pattern):
    kwargs = {}
    for param, converter in url_pattern.pattern.converters.items():
        conv_name = _CONVERTER_NAME_BY_CLASS.get(type(converter), "str")
        kwargs[param] = _DUMMY_VALUE_BY_CONVERTER.get(conv_name, "dummy")
    return kwargs


def _api_key_gated(url_pattern):
    """
    Whether this route's view is gated by rest_framework_api_key's HasAPIKey
    rather than (or as well as) a session.

    DRF's as_view() stashes the view class on the function it returns as
    `.cls`, and functools.wraps -- which page_access_required and friends in
    insurance/urls.py all use -- copies the wrapped function's __dict__ onto
    the wrapper. So this still finds the class through a session wrapper,
    which matters for api-export-rates: it stacks page_access_required on top
    of a HasAPIKey view.
    """
    view_class = getattr(url_pattern.callback, "cls", None)
    return HasAPIKey in (getattr(view_class, "permission_classes", None) or ())


class UrlAuthGateTests(TestCase):
    """
    Every page in insurance/urls.py must be gated behind login_required (or
    page_access_required/staff_required, which both wrap it) unless it is
    explicitly whitelisted in PUBLIC_URL_NAMES above. An anonymous GET to a
    gated URL must redirect to the login page, never return 200 or any
    other status.

    The one exception is a server-to-server route gated by HasAPIKey: those
    refuse a keyless request with 403 rather than redirecting a browser, so
    they are held to "403 or a 302 to login" instead. They are detected
    directly rather than whitelisted -- see _api_key_gated.
    """

    def test_every_named_url_requires_login_unless_whitelisted(self):
        client = Client()
        login_path = reverse("login")
        checked = 0
        # Collected rather than asserted inline: a bare assert stops at the
        # first bad route and hides every one after it in urlpatterns order,
        # which is exactly how three unlisted HasAPIKey routes sat unnoticed
        # behind a fourth. One failure should report the whole list.
        failures = []

        for pattern in insurance_urls.urlpatterns:
            name = getattr(pattern, "name", None)
            if not name or name in PUBLIC_URL_NAMES:
                continue

            path = reverse(name, kwargs=_dummy_kwargs(pattern))
            response = client.get(path)
            checked += 1
            redirects_to_login = (
                response.status_code == 302 and response.url.startswith(login_path)
            )

            if _api_key_gated(pattern):
                # Server-to-server route: rest_framework_api_key refuses a
                # keyless request with 403 rather than redirecting a browser.
                # A 302 to login is equally fine — api-export-rates stacks
                # page_access_required on top, and that fires first. Either
                # way the route is not reachable anonymously, which is what
                # this test is actually guarding. Anything else (a 200 above
                # all, i.e. DRF's AllowAny default) is a real hole.
                if response.status_code != 403 and not redirects_to_login:
                    failures.append(
                        f"'{name}' ({path}) is HasAPIKey-gated but returned "
                        f"{response.status_code} for an anonymous, keyless request — "
                        f"expected 403 from the API-key check, or a 302 to login from "
                        f"a session wrapper stacked on top."
                    )
                continue

            if response.status_code != 302:
                # A DRF route lands here when it is neither HasAPIKey-gated nor
                # session-wrapped. DEFAULT_PERMISSION_CLASSES is IsAuthenticated
                # (project/settings.py), so it is not *open* — but this app has
                # no such route by design, and the choice between the two gates
                # should be made deliberately rather than inherited, so say that
                # instead of sending someone hunting for a missing wrapper.
                if getattr(pattern.callback, "cls", None) is not None:
                    failures.append(
                        f"'{name}' ({path}) is a DRF view that is neither HasAPIKey-gated "
                        f"nor wrapped in insurance/urls.py, so it returned "
                        f"{response.status_code} instead of redirecting to login. Add "
                        f"permission_classes = [HasAPIKey] if it is server-to-server, or "
                        f"a page_access_required wrapper if a browser session should reach it."
                    )
                else:
                    failures.append(
                        f"'{name}' ({path}) returned {response.status_code} for an "
                        f"anonymous request instead of redirecting to login — it may "
                        f"be missing a login_required/page_access_required/"
                        f"staff_required wrapper in insurance/urls.py."
                    )
            elif not redirects_to_login:
                failures.append(
                    f"'{name}' ({path}) redirected an anonymous request to "
                    f"'{response.url}' instead of the login page."
                )

        self.assertFalse(
            failures,
            f"{len(failures)} route(s) did not refuse an anonymous request:\n  "
            + "\n  ".join(failures)
        )

        # Guards against this test silently checking nothing if the loop
        # above ever stops matching real entries in insurance/urls.py.
        self.assertGreater(
            checked, 20,
            "Expected to check most of insurance/urls.py's named routes — "
            "did the import path or urlpatterns structure change?"
        )

    def test_whitelisted_public_urls_still_exist(self):
        """
        Sanity check the other direction: every name in PUBLIC_URL_NAMES
        must still be a real route. Catches a stale whitelist entry left
        behind after a route is renamed or removed.
        """
        actual_names = {p.name for p in insurance_urls.urlpatterns if getattr(p, "name", None)}
        stale = PUBLIC_URL_NAMES - actual_names
        self.assertFalse(stale, f"Whitelisted public URL name(s) no longer exist: {stale}")


class ApiKeyGatedRouteTests(TestCase):
    """
    The server-to-server routes UrlAuthGateTests above hands off to
    _api_key_gated. That handoff is only safe if the detection actually works
    and the gate actually refuses, so both are asserted here rather than
    assumed — otherwise a broken detector would quietly turn the auth gate
    test into a rubber stamp for every API route.
    """

    def setUp(self):
        self.client = Client()
        self.api_key_routes = [
            p for p in insurance_urls.urlpatterns
            if getattr(p, "name", None) and _api_key_gated(p)
        ]

    def test_the_api_key_routes_are_found(self):
        # If DRF ever stops setting `.cls` on as_view()'s return value, or a
        # wrapper stops using functools.wraps, this drops to zero -- and every
        # one of those routes would then be held to the strict 302 rule, which
        # fails loudly rather than silently passing. This just names the
        # breakage directly instead of leaving it to be inferred.
        found = {p.name for p in self.api_key_routes}
        self.assertEqual(
            found,
            {
                "api-export-rates",
                "api_policy_lock_checker",
                "api_health_payout_rates",
                "api_make_model_master",
                "sso_issue_ticket",
            },
            "The set of HasAPIKey-gated routes changed. If that was deliberate, update "
            "this list; if not, _api_key_gated may have stopped seeing through as_view() "
            "or a urls.py wrapper.",
        )

    def test_detection_sees_through_a_session_wrapper(self):
        # api-export-rates is the awkward one: page_access_required stacked on
        # top of a HasAPIKey view. functools.wraps copies as_view()'s `.cls`
        # onto the wrapper, which is the only reason this works.
        wrapped = next(p for p in insurance_urls.urlpatterns if p.name == "api-export-rates")
        self.assertTrue(_api_key_gated(wrapped))
        self.assertIsNot(
            wrapped.callback, wrapped.callback.cls,
            "Expected api-export-rates to be a wrapper around the DRF view, not the view itself.",
        )

    def test_every_api_key_route_refuses_a_keyless_request(self):
        refused = []
        for pattern in self.api_key_routes:
            path = reverse(pattern.name, kwargs=_dummy_kwargs(pattern))
            for method in (self.client.get, self.client.post):
                response = method(path)
                self.assertNotIn(
                    response.status_code, (200, 201),
                    f"'{pattern.name}' ({path}) served a keyless "
                    f"{method.__name__.upper()} with {response.status_code}.",
                )
            refused.append(pattern.name)
        self.assertEqual(len(refused), len(self.api_key_routes))

    def test_a_valid_api_key_gets_past_the_gate(self):
        # The mirror image of the test above: proves the 403s there come from
        # the missing key, not from the route being broken for everyone.
        _, key = APIKey.objects.create_key(name="test-gate-probe")
        pattern = next(p for p in self.api_key_routes if p.name == "api_make_model_master")
        path = reverse(pattern.name)

        self.assertIn(self.client.get(path).status_code, (401, 403))
        self.assertEqual(
            self.client.get(path, HTTP_AUTHORIZATION=f"Api-Key {key}").status_code, 200
        )


@override_settings(
    PARTNER_SSO_TICKET_SECRET="test-secret-only-used-by-this-test-class",
    # DatabaseCache (the real backend, see project/settings.py) needs
    # `manage.py createcachetable`, which isn't run against the test
    # database -- swap in an in-memory cache so the single-use check below
    # exercises real cache behavior instead of failing closed for an
    # unrelated reason.
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
)
class SSOTicketTests(TestCase):
    """Unit tests for insurance/sso.py's partner-portal ticket mint/verify round trip."""

    def setUp(self):
        self.user = User.objects.create_user(username="partner.user@example.com")

    def test_valid_ticket_round_trips_once(self):
        ticket = sso.mint_ticket(self.user, "policy_lock_checker", "jti-1")
        result = sso.verify_and_consume_ticket(ticket)
        self.assertIsNotNone(result)
        user, landing_page = result
        self.assertEqual(user, self.user)
        self.assertEqual(landing_page, "policy_lock_checker")

    def test_replayed_ticket_is_rejected(self):
        ticket = sso.mint_ticket(self.user, "policy_lock_checker", "jti-2")
        self.assertIsNotNone(sso.verify_and_consume_ticket(ticket))
        # Second use of the exact same ticket must fail even though it
        # hasn't expired yet.
        self.assertIsNone(sso.verify_and_consume_ticket(ticket))

    def test_expired_ticket_is_rejected(self):
        with mock.patch("insurance.sso.time.time", return_value=1_000_000.0):
            ticket = sso.mint_ticket(self.user, "policy_lock_checker", "jti-3")
        with mock.patch("insurance.sso.time.time", return_value=1_000_000.0 + sso.TICKET_TTL_SECONDS + 1):
            self.assertIsNone(sso.verify_and_consume_ticket(ticket))

    def test_tampered_signature_is_rejected(self):
        ticket = sso.mint_ticket(self.user, "policy_lock_checker", "jti-4")
        payload, _, _sig = ticket.rpartition(".")
        forged = f"{payload}.{'0' * 64}"
        self.assertIsNone(sso.verify_and_consume_ticket(forged))

    def test_wrong_audience_is_rejected(self):
        payload = f"{self.user.username}:{int(sso.time.time()) + 60}:policy_lock_checker:some-other-audience:jti-5"
        forged = f"{payload}.{sso._sign(payload)}"
        self.assertIsNone(sso.verify_and_consume_ticket(forged))

    def test_unknown_user_is_rejected(self):
        payload = f"nobody@example.com:{int(sso.time.time()) + 60}:policy_lock_checker:{sso.AUDIENCE}:jti-6"
        forged = f"{payload}.{sso._sign(payload)}"
        self.assertIsNone(sso.verify_and_consume_ticket(forged))


@override_settings(
    PARTNER_SSO_TICKET_SECRET="test-secret-only-used-by-this-test-class",
    # See the comment on SSOTicketTests above -- DatabaseCache needs a table
    # this test database doesn't have.
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
    # ManifestStaticFilesStorage (the real "staticfiles" backend, see
    # project/settings.py) needs `manage.py collectstatic` to have been run
    # to generate its manifest -- not something a fresh test run can rely
    # on. Swap in the plain (non-hashed) storage just for this class, which
    # actually renders full pages (policy_lock_checker.html, via base.html).
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    },
)
class SSOHandoffViewTests(TestCase):
    """
    End-to-end tests for the two views that make up the partner-portal SSO
    handoff: IssueSSOTicketAPIView (server-to-server) and sso_consume_view
    (the browser-redirect landing page).
    """

    def setUp(self):
        self.client = Client()
        _, self.api_key = APIKey.objects.create_key(name="test-partner")
        for name in ("Can_View_Policy_Locker", "Can_View_Health_Payout_Rates", "Can_View_Life_Payout_Grid"):
            Group.objects.get_or_create(name=name)

    def _issue(self, **overrides):
        payload = {
            "email": "agent@arhamsecure.com",
            "full_name": "Test Agent",
            "requested_pages": ["Rate_Checker"],
        }
        payload.update(overrides)
        return self.client.post(
            reverse("sso_issue_ticket"),
            data=json.dumps(payload),
            content_type="application/json",
            HTTP_AUTHORIZATION=f"Api-Key {self.api_key}",
        )

    @staticmethod
    def _path_from_redirect_url(redirect_url):
        # redirect_url is absolute (http://host/sso/consume/?ticket=...) --
        # the test client wants a path.
        return "/" + redirect_url.split("://", 1)[1].split("/", 1)[1]

    def test_issue_ticket_requires_api_key(self):
        response = self.client.post(
            reverse("sso_issue_ticket"),
            data=json.dumps({"email": "agent@arhamsecure.com"}),
            content_type="application/json",
        )
        self.assertIn(response.status_code, (401, 403))

    def test_full_handoff_logs_user_in_with_scoped_access(self):
        response = self._issue()
        self.assertEqual(response.status_code, 200)

        user = User.objects.get(username="agent@arhamsecure.com")
        self.assertTrue(user.is_active)
        self.assertFalse(user.has_usable_password())
        self.assertEqual(
            set(user.groups.values_list("name", flat=True)),
            {"Can_View_Policy_Locker", "Can_View_Health_Payout_Rates", "Can_View_Life_Payout_Grid"},
        )

        # Default landing page is the Rate Checker hub, not a specific tab...
        consume_path = self._path_from_redirect_url(response.json()["redirect_url"])
        consume_response = self.client.get(consume_path)
        self.assertEqual(consume_response.status_code, 302)
        self.assertEqual(consume_response.url, reverse("rate_checker"))

        # ...which itself redirects to the highest-priority accessible tab (Motor).
        hub_response = self.client.get(reverse("rate_checker"))
        self.assertEqual(hub_response.url, reverse("policy_lock_checker"))

        # Session is now live: a granted page loads directly...
        self.assertEqual(self.client.get(reverse("policy_lock_checker")).status_code, 200)
        self.assertEqual(self.client.get(reverse("health_payout_rates")).status_code, 200)
        # ...but a page outside the granted set is still forbidden.
        self.assertEqual(self.client.get(reverse("audit_logs")).status_code, 403)

    def test_explicit_landing_page_still_lands_directly_on_a_tab(self):
        response = self._issue(landing_page="health_payout_rates")
        consume_path = self._path_from_redirect_url(response.json()["redirect_url"])
        consume_response = self.client.get(consume_path)
        self.assertEqual(consume_response.url, reverse("health_payout_rates"))

    def test_ticket_cannot_be_replayed(self):
        response = self._issue()
        consume_path = self._path_from_redirect_url(response.json()["redirect_url"])

        first = self.client.get(consume_path)
        self.assertEqual(first.url, reverse("rate_checker"))

        self.client.logout()
        second = self.client.get(consume_path)
        self.assertEqual(second.url, reverse("login"))

    def test_scope_is_clamped_server_side(self):
        response = self._issue(requested_pages=["Can_View_Policy_Locker", "ADMIN", "Can_View_Audit_Log"])
        self.assertEqual(response.status_code, 200)
        user = User.objects.get(username="agent@arhamsecure.com")
        self.assertEqual(set(user.groups.values_list("name", flat=True)), {"Can_View_Policy_Locker"})

    def test_previously_grantable_pages_no_longer_granted(self):
        """
        Rate Checker replaced the old allow-list (Policy Locker + Locked
        Policies + Motor Points Logs from the first version of this
        handoff) -- Locked Policies / Motor Points Logs are no longer
        requestable through this integration.
        """
        response = self._issue(requested_pages=[
            "Can_View_Policy_Locker", "Can_View_Locked_Policies", "Can_View_Motor_Points_Logs",
        ])
        self.assertEqual(response.status_code, 200)
        user = User.objects.get(username="agent@arhamsecure.com")
        self.assertEqual(set(user.groups.values_list("name", flat=True)), {"Can_View_Policy_Locker"})

    def test_direct_visit_with_no_ticket_falls_back_to_login(self):
        response = self.client.get(reverse("sso_consume"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("login"))


@override_settings(
    # See the comment on SSOHandoffViewTests above -- the 403 page extends
    # base.html, which needs this when ManifestStaticFilesStorage's
    # manifest hasn't been generated (no `collectstatic` in this test run).
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    },
)
class RateCheckerEntryTests(TestCase):
    """Tests for the plain (non-SSO) Rate Checker sidebar entry point."""

    def setUp(self):
        self.client = Client()
        for name in ("Can_View_Policy_Locker", "Can_View_Health_Payout_Rates", "Can_View_Life_Payout_Grid"):
            Group.objects.get_or_create(name=name)
        self.user = User.objects.create_user(username="broker", password="a-strong-test-password-1")
        self.client.force_login(self.user)

    def test_redirects_to_highest_priority_accessible_tab(self):
        self.user.groups.add(Group.objects.get(name="Can_View_Health_Payout_Rates"))
        self.user.groups.add(Group.objects.get(name="Can_View_Life_Payout_Grid"))
        response = self.client.get(reverse("rate_checker"))
        # Motor isn't granted, so Health (next in priority) wins.
        self.assertEqual(response.url, reverse("health_payout_rates"))

    def test_falls_through_to_life_if_thats_all_thats_granted(self):
        self.user.groups.add(Group.objects.get(name="Can_View_Life_Payout_Grid"))
        response = self.client.get(reverse("rate_checker"))
        self.assertEqual(response.url, reverse("life_payout_grid_redirect"))

    def test_forbidden_with_none_of_the_three_groups(self):
        response = self.client.get(reverse("rate_checker"))
        self.assertEqual(response.status_code, 403)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class GridSummaryDateFilterTests(TestCase):
    """Rate Master Health > Grid Summary's "Valid As Of" filter."""

    def setUp(self):
        from datetime import date
        from insurance.models import ProductMaster, RateGroup, RateMaster

        self.client = Client()
        Group.objects.get_or_create(name="Can_View_Rate_Master_Health")
        self.user = User.objects.create_user(username="ops", password="a-strong-test-password-1")
        self.user.groups.add(Group.objects.get(name="Can_View_Rate_Master_Health"))
        self.client.force_login(self.user)

        product = ProductMaster.objects.create(name="Private Car")
        RateMaster.objects.create(
            insurance_company="Acme General", product=product, status="ACTIVE", is_deleted="NO",
            group=RateGroup.objects.create(key_hash="h1"),
            from_date=date(2026, 1, 1), to_date=date(2026, 6, 30),
        )
        RateMaster.objects.create(
            insurance_company="Zenith Insurance", product=product, status="ACTIVE", is_deleted="NO",
            group=RateGroup.objects.create(key_hash="h2"),
            from_date=date(2026, 7, 1), to_date=date(2026, 12, 31),
        )

    def test_no_filter_shows_every_grid(self):
        response = self.client.get(reverse("rate_master_health"), {"view": "grid"})
        self.assertEqual(response.context["grid_summary_total_combinations"], 2)

    def test_soft_deleted_rows_are_excluded_from_the_pivot(self):
        # is_deleted="YES" is the Rate Master dashboard's own "Is Deleted?"
        # flag - a soft-deleted grid must not still count here as uploaded.
        from datetime import date
        from insurance.models import ProductMaster, RateGroup, RateMaster
        RateMaster.objects.create(
            insurance_company="Deleted Co", product=ProductMaster.objects.create(name="Two Wheeler"),
            status="ACTIVE", is_deleted="YES", group=RateGroup.objects.create(key_hash="h3"),
            from_date=date(2026, 1, 1), to_date=date(2026, 6, 30),
        )
        response = self.client.get(reverse("rate_master_health"), {"view": "grid"})
        self.assertEqual(response.context["grid_summary_total_combinations"], 2)
        rows = list(response.context["grid_summary_page_obj"])
        self.assertFalse(any(r["insurance_company"] == "Deleted Co" for r in rows))

    def test_inactive_but_not_deleted_rows_are_still_included(self):
        # Deliberately different from is_deleted: an insurer's grid going
        # INACTIVE is a real, current fact worth seeing - only a soft
        # DELETE removes it from the pivot.
        from datetime import date
        from insurance.models import ProductMaster, RateGroup, RateMaster
        RateMaster.objects.create(
            insurance_company="Inactive Co", product=ProductMaster.objects.create(name="Two Wheeler"),
            status="INACTIVE", is_deleted="NO", group=RateGroup.objects.create(key_hash="h4"),
            from_date=date(2026, 1, 1), to_date=date(2026, 6, 30),
        )
        response = self.client.get(reverse("rate_master_health"), {"view": "grid"})
        self.assertEqual(response.context["grid_summary_total_combinations"], 3)
        rows = list(response.context["grid_summary_page_obj"])
        self.assertTrue(any(r["insurance_company"] == "Inactive Co" for r in rows))

    def test_as_of_date_narrows_to_grids_valid_on_that_date(self):
        response = self.client.get(
            reverse("rate_master_health"), {"view": "grid", "grid_as_of_date": "2026-03-15"}
        )
        rows = list(response.context["grid_summary_page_obj"])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["insurance_company"], "Acme General")

    def test_an_invalid_date_is_ignored_rather_than_erroring(self):
        response = self.client.get(
            reverse("rate_master_health"), {"view": "grid", "grid_as_of_date": "not-a-date"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["grid_summary_total_combinations"], 2)

    def test_export_honours_the_same_filter(self):
        response = self.client.get(
            reverse("export_grid_summary_xlsx"), {"grid_as_of_date": "2026-03-15"}
        )
        self.assertEqual(response.status_code, 200)

        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 1)
        self.assertEqual(data_rows[0][2], "Acme General")

    def test_export_with_no_filter_exports_everything(self):
        response = self.client.get(reverse("export_grid_summary_xlsx"))
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 2)


class RangeFilterInvalidInputTests(TestCase):
    """
    apply_range_filter() (shared by Rate Master's dashboard/export and Health
    Rate Master) feeds request.GET values straight into a filter on a
    FloatField. A value Django can't coerce with float() -- e.g. "12/" or
    "12\\", typed into Age Range / CC Range / SC Range -- used to raise an
    uncaught ValueError there and 500 the page, instead of being ignored the
    way an invalid date already is elsewhere on this same dashboard (see
    GridSummaryDateFilterTests.test_an_invalid_date_is_ignored_rather_than_erroring).
    """

    def setUp(self):
        from insurance.models import ProductMaster, RateGroup, RateMaster

        self.client = Client()
        Group.objects.get_or_create(name="Can_View_Dashboard")
        Group.objects.get_or_create(name="Can_View_Health_Rate_Master")
        self.user = User.objects.create_user(username="ops", password="a-strong-test-password-1")
        self.user.groups.add(
            Group.objects.get(name="Can_View_Dashboard"),
            Group.objects.get(name="Can_View_Health_Rate_Master"),
        )
        self.client.force_login(self.user)

        product = ProductMaster.objects.create(name="Private Car")
        RateMaster.objects.create(
            insurance_company="Acme General", product=product, status="ACTIVE", is_deleted="NO",
            group=RateGroup.objects.create(key_hash="range-filter-h1"),
            vehicle_age_min=0, vehicle_age_max=5,
        )

    def test_slash_or_backslash_in_age_range_is_ignored_rather_than_erroring(self):
        # dashboard() only runs its real query once >=2 filters are set --
        # match that here with a harmless second filter (status).
        for bad_value in ["12/", "12\\"]:
            response = self.client.get(reverse("dashboard"), {"age_range": bad_value, "status": "ACTIVE"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.context["total"], 1)

    def test_slash_in_cc_or_sc_range_is_ignored_rather_than_erroring(self):
        response = self.client.get(
            reverse("dashboard"), {"cc_range": "12/", "sc_range": "3\\", "status": "ACTIVE"}
        )
        self.assertEqual(response.status_code, 200)

    def test_export_honours_the_same_graceful_handling(self):
        response = self.client.get(reverse("export_rates_xlsx"), {"age_range": "12/"})
        self.assertEqual(response.status_code, 200)

    def test_health_rate_master_range_filters_are_also_protected(self):
        response = self.client.get(reverse("health_rate_master"), {
            "age_range": "12/", "sum_insured_range": "1\\", "deductible_range": "x",
        })
        self.assertEqual(response.status_code, 200)

    def test_valid_bare_age_range_still_filters(self):
        # 5 falls inside the fixture row's [0, 5] span -> still matches.
        response = self.client.get(reverse("dashboard"), {"age_range": "5", "status": "ACTIVE"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total"], 1)

    def test_out_of_range_age_still_excludes(self):
        # Proves invalid input is ignored, not that filtering stopped working.
        response = self.client.get(reverse("dashboard"), {"age_range": "99", "status": "ACTIVE"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total"], 0)


class ApiUploadChunkRateMasterDedupTests(TestCase):
    """
    Bulk Upload -> Rate Master (api_upload_chunk) used to insert a fresh
    RateMaster row for every row it was handed with no check for an
    identical row already existing in the same group -- so a source file
    that listed the same rate line twice, the browser resending an
    already-saved chunk, or a whole file getting re-uploaded all produced
    exact-duplicate rows sharing one group_id. Confirmed in production:
    108,516 duplicate rows across 21,532 groups (~14% of the table).

    The fix makes the insert idempotent per (group, new_rto_list) -- see
    existing_rto_by_group in api_upload_chunk's rate_master branch.
    """

    def setUp(self):
        from insurance.models import RTOMaster

        self.client = Client()
        Group.objects.get_or_create(name="Can_Upload_CSV")
        self.user = User.objects.create_user(username="uploader", password="a-strong-test-password-1")
        self.user.groups.add(Group.objects.get(name="Can_Upload_CSV"))
        self.client.force_login(self.user)

        RTOMaster.objects.create(rto_name="MUMBAI")
        RTOMaster.objects.create(rto_name="PUNE")

    def _row(self, rto="MUMBAI", **overrides):
        row = {
            "insurance_company": "Acme General",
            "new_rto_list": rto,
            "from_date": "2026-01-01",
            "to_date": "2026-12-31",
            "pi_od_rate": "10",
        }
        row.update(overrides)
        return row

    def _upload(self, rows, upload_batch_id="batch-1"):
        return self.client.post(
            reverse("api_upload_chunk"),
            data=json.dumps({
                "target_table": "rate_master",
                "rows": rows,
                "upload_batch_id": upload_batch_id,
                "dry_run": False,
            }),
            content_type="application/json",
        )

    def test_same_row_repeated_in_one_chunk_is_only_saved_once(self):
        from insurance.models import RateMaster

        response = self._upload([self._row(), self._row()])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(RateMaster.objects.count(), 1)

    def test_same_chunk_resent_in_a_second_request_is_not_duplicated(self):
        # Simulates a browser/network retry resending a chunk that already
        # succeeded, or the whole file being uploaded twice -- same
        # upload_batch_id both times, since that's what a single Upload
        # click reuses across its chunks/passes (see upload.html).
        from insurance.models import RateMaster

        first = self._upload([self._row()])
        second = self._upload([self._row()])
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(RateMaster.objects.count(), 1)

    def test_genuinely_different_rtos_in_the_same_group_both_saved(self):
        # Guards against the dedup check being too broad: two rows that
        # differ only by new_rto_list are still two legitimate rows.
        from insurance.models import RateMaster

        response = self._upload([self._row(rto="MUMBAI"), self._row(rto="PUNE")])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(RateMaster.objects.count(), 2)

    def test_reuploading_content_that_is_already_active_does_not_create_a_duplicate(self):
        # A *different* upload_batch_id simulates a genuinely separate upload
        # session (e.g. exporting an existing grid and re-importing it
        # unchanged) -- not a retry of the same one. Uploads always insert as
        # INACTIVE, so "already active" here means some later, separate step
        # turned this content live; re-uploading it shouldn't add a parallel
        # copy of a rate that's already in effect.
        from insurance.models import RateMaster

        self._upload([self._row()], upload_batch_id="original-upload")
        RateMaster.objects.update(status="ACTIVE")

        response = self._upload([self._row()], upload_batch_id="a-completely-different-upload")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(RateMaster.objects.count(), 1)

    def test_reuploading_when_existing_row_is_inactive_still_creates_a_new_one(self):
        # The guard only matches ACTIVE, non-deleted rows -- content that was
        # never activated (or was deliberately turned off) doesn't block a
        # legitimate re-upload of that same content.
        from insurance.models import RateMaster

        self._upload([self._row()], upload_batch_id="original-upload")
        self.assertEqual(RateMaster.objects.filter(status="INACTIVE").count(), 1)

        response = self._upload([self._row()], upload_batch_id="a-completely-different-upload")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(RateMaster.objects.count(), 2)


class DedupeRateMasterCommandTests(TestCase):
    """
    dedupe_rate_master soft-deletes the exact-duplicate rows produced by the
    now-fixed api_upload_chunk bug, keeping the lowest id per duplicate
    cluster. Manual testing surfaced a real bug here: the duplicate
    signature originally included is_deleted, so every already-soft-deleted
    row looked like a fresh duplicate of every other already-soft-deleted
    row on the very next scan (they're identical except id) -- re-running
    the command after a cleanup pass kept finding "more" duplicates
    forever. Fixed by scoping detection to is_deleted="NO" rows only.
    """

    def setUp(self):
        from datetime import date
        from insurance.models import ProductMaster, RateGroup, RateMaster

        product = ProductMaster.objects.create(name="Private Car")
        group = RateGroup.objects.create(key_hash="dedupe-test-group")

        def make(rto):
            return RateMaster.objects.create(
                group=group, product=product, insurance_company="Acme General",
                status="ACTIVE", is_deleted="NO", new_rto_list=rto,
                from_date=date(2026, 1, 1), to_date=date(2026, 12, 31),
            )

        # Group's real content: MUMBAI once, PUNE three times over (2 extras).
        self.kept_mumbai = make("MUMBAI")
        self.kept_pune = make("PUNE")
        self.dup_pune_1 = make("PUNE")
        self.dup_pune_2 = make("PUNE")

    def test_dry_run_reports_but_does_not_modify(self):
        from io import StringIO
        from django.core.management import call_command
        from insurance.models import RateMaster

        out = StringIO()
        call_command("dedupe_rate_master", stdout=out)
        self.assertIn("2", out.getvalue())
        self.assertEqual(RateMaster.objects.filter(is_deleted="YES").count(), 0)

    def test_apply_keeps_lowest_id_per_cluster_and_logs_it(self):
        from django.core.management import call_command
        from insurance.models import AuditLog, RateMaster

        call_command("dedupe_rate_master", apply=True)

        self.kept_mumbai.refresh_from_db()
        self.kept_pune.refresh_from_db()
        self.dup_pune_1.refresh_from_db()
        self.dup_pune_2.refresh_from_db()

        self.assertEqual(self.kept_mumbai.is_deleted, "NO")
        self.assertEqual(self.kept_pune.is_deleted, "NO")
        self.assertEqual(self.dup_pune_1.is_deleted, "YES")
        self.assertEqual(self.dup_pune_2.is_deleted, "YES")

        self.assertTrue(AuditLog.objects.filter(action="BULK DEDUPE").exists())

    def test_rerunning_after_apply_finds_nothing_left(self):
        from io import StringIO
        from django.core.management import call_command

        call_command("dedupe_rate_master", apply=True)

        out = StringIO()
        call_command("dedupe_rate_master", stdout=out)
        self.assertIn("Nothing to do.", out.getvalue())


class CrossGroupDedupeTests(TestCase):
    """
    --cross-group is the second, separate duplicate category found after
    the first cleanup shipped: identical-content rows living under
    *different* group_ids (e.g. exporting a grid and re-importing it
    unchanged creates a brand new group with the same content as an
    existing one). Default-mode dedupe_rate_master never compares across
    groups, so it can't see these -- only --cross-group can.
    """

    def setUp(self):
        from datetime import date
        from insurance.models import ProductMaster, RateGroup, RateMaster

        self.product = ProductMaster.objects.create(name="Private Car")

        def make(group, status):
            return RateMaster.objects.create(
                group=group, product=self.product, insurance_company="Acme General",
                status=status, is_deleted="NO", new_rto_list="MUMBAI",
                from_date=date(2026, 1, 1), to_date=date(2026, 12, 31),
            )

        # Same exact content, two different groups -- as if re-uploaded later.
        self.original = make(RateGroup.objects.create(key_hash="original-group"), "ACTIVE")
        self.reupload = make(RateGroup.objects.create(key_hash="reupload-group"), "ACTIVE")

    def test_default_mode_does_not_see_cross_group_duplicates(self):
        from io import StringIO
        from django.core.management import call_command

        out = StringIO()
        call_command("dedupe_rate_master", stdout=out)
        self.assertIn("Nothing to do.", out.getvalue())

    def test_cross_group_mode_keeps_lowest_id_and_soft_deletes_the_rest(self):
        from django.core.management import call_command
        from insurance.models import RateMaster

        call_command("dedupe_rate_master", apply=True, cross_group=True)

        self.original.refresh_from_db()
        self.reupload.refresh_from_db()
        self.assertEqual(self.original.is_deleted, "NO")
        self.assertEqual(self.reupload.is_deleted, "YES")

    def test_cross_group_mode_ignores_inactive_matches(self):
        # Mirrors the upload-time guard: only ACTIVE rows count as "already
        # exists" for this sweep too, so an inactive row with matching
        # content is left alone rather than being silently removed.
        from django.core.management import call_command
        from insurance.models import RateGroup, RateMaster

        self.reupload.status = "INACTIVE"
        self.reupload.save()

        call_command("dedupe_rate_master", apply=True, cross_group=True)

        self.reupload.refresh_from_db()
        self.assertEqual(self.reupload.is_deleted, "NO")


class PurgeDeletedRateMasterCommandTests(TestCase):
    """
    purge_deleted_rate_master PERMANENTLY deletes every is_deleted=YES row
    -- unlike dedupe_rate_master, this is a real DELETE, not a flag flip,
    and it's scoped to *every* soft-deleted row table-wide (not just ones
    from a specific cleanup run) per explicit instruction, since is_deleted
    is also a normal, independently-used feature of the app.
    """

    def setUp(self):
        from datetime import date
        from insurance.models import ProductMaster, RateGroup, RateMaster

        product = ProductMaster.objects.create(name="Private Car")
        group = RateGroup.objects.create(key_hash="purge-test-group")

        def make(is_deleted):
            return RateMaster.objects.create(
                group=group, product=product, insurance_company="Acme General",
                status="ACTIVE", is_deleted=is_deleted, new_rto_list="MUMBAI",
                from_date=date(2026, 1, 1), to_date=date(2026, 12, 31),
            )

        self.kept_row = make("NO")
        self.deleted_row_1 = make("YES")
        self.deleted_row_2 = make("YES")

    def test_dry_run_reports_but_does_not_delete(self):
        from io import StringIO
        from django.core.management import call_command
        from insurance.models import RateMaster

        out = StringIO()
        call_command("purge_deleted_rate_master", stdout=out)
        self.assertIn("2", out.getvalue())
        self.assertEqual(RateMaster.objects.count(), 3)

    def test_apply_permanently_deletes_only_is_deleted_yes_rows(self):
        from django.core.management import call_command
        from insurance.models import AuditLog, RateMaster

        call_command("purge_deleted_rate_master", apply=True)

        remaining_ids = set(RateMaster.objects.values_list("id", flat=True))
        self.assertEqual(remaining_ids, {self.kept_row.id})
        self.assertTrue(AuditLog.objects.filter(action="BULK PERMANENT DELETE").exists())

    def test_locked_policy_referencing_a_deleted_row_is_not_broken(self):
        from django.core.management import call_command
        from insurance.models import LockedPolicy

        lock = LockedPolicy.objects.create(
            source_rate=self.deleted_row_1, vehicle_no="MH12AB1234", policy_holder_name="Test Holder",
        )

        call_command("purge_deleted_rate_master", apply=True)

        lock.refresh_from_db()
        self.assertIsNone(lock.source_rate)


class RateDecimalRoundingTests(TestCase):
    """
    Dashboard rate fields (pi_od_rate, pi_tp_rate, pi_tp_2..5, pi_net_rate,
    pi_flat_amount, pi_vli, po_od_rate, po_tp_rate, po_net_rate,
    po_flat_amount) must never be saved with more than 2 decimal places,
    whether set via the single-record Edit form or Bulk Update Selected.
    """

    def setUp(self):
        from insurance.models import ProductMaster, RateGroup, RateMaster

        self.client = Client()
        Group.objects.get_or_create(name="Can_View_Dashboard")
        self.user = User.objects.create_user(username="ops", password="a-strong-test-password-1")
        self.user.groups.add(Group.objects.get(name="Can_View_Dashboard"))
        self.client.force_login(self.user)

        self.product = ProductMaster.objects.create(name="Private Car")
        self.group = RateGroup.objects.create(key_hash="decimal-rounding-group")
        self.record = RateMaster.objects.create(
            group=self.group, product=self.product, insurance_company="Acme General",
            status="ACTIVE", is_deleted="NO", pi_od_rate=10.0, po_od_rate=3.0,
        )

    def test_bulk_update_rounds_a_targeted_rate_field_to_two_decimals(self):
        from insurance.models import RateMaster

        self.client.post(reverse("bulk_update_rates"), {
            "selected_groups": str(self.group.id),
            "update_field": "pi_od_rate",
            "update_value": "12.34567",
        })
        self.record.refresh_from_db()
        self.assertEqual(self.record.pi_od_rate, 12.35)

    def test_bulk_update_does_not_round_unrelated_numeric_fields(self):
        # tariff_min/max, cc_min/max, vehicle_age_min/max, sc_min/max are
        # numeric too but were never in the two-decimal-place requirement.
        from insurance.models import RateMaster

        self.client.post(reverse("bulk_update_rates"), {
            "selected_groups": str(self.group.id),
            "update_field": "tariff_min",
            "update_value": "12.34567",
        })
        self.record.refresh_from_db()
        self.assertEqual(self.record.tariff_min, 12.34567)

    def test_edit_form_rounds_rate_field_to_two_decimals(self):
        from insurance.views import RateForm

        form = RateForm(data={
            "insurance_company": "Acme General",
            "status": "ACTIVE",
            "is_deleted": "NO",
            "pi_od_rate": "12.3456",
            "po_od_rate": "3.005",
        }, instance=self.record, initial={"new_rto_list": [], "new_vehicle_makes": []})

        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["pi_od_rate"], 12.35)
        self.assertEqual(form.cleaned_data["po_od_rate"], 3.0)


class PiPoMarginFloatPrecisionTests(TestCase):
    """
    Rate Master Health's "Pi vs Po Validation Errors" cards
    (_rate_master_pi_po_rate_violations_qs) flag a row when Pi - Po isn't
    exactly RATE_MASTER_PI_PO_MARGIN (7). pi_od_rate/po_od_rate etc. are
    plain FloatFields (IEEE-754 doubles), so e.g. 19.65 - 12.65 computes to
    6.999999999999998 in raw floating point rather than 7.0 -- a
    mathematically exact margin was being flagged as a violation purely
    from binary float representation error.
    """

    def setUp(self):
        from insurance.models import ProductMaster, RateGroup, RateMaster

        product = ProductMaster.objects.create(name="Private Car")
        group = RateGroup.objects.create(key_hash="pi-po-float-precision-group")

        # 19.65 - 12.65 == 7 mathematically, but not in raw IEEE-754 double
        # arithmetic -- this is the exact pair reported as a false positive.
        self.exact_margin_row = RateMaster.objects.create(
            group=group, product=product, insurance_company="Acme General",
            status="ACTIVE", is_deleted="NO", pi_od_rate=19.65, po_od_rate=12.65,
        )
        self.genuine_violation_row = RateMaster.objects.create(
            group=group, product=product, insurance_company="Acme General",
            status="ACTIVE", is_deleted="NO", pi_od_rate=20.0, po_od_rate=12.0,
        )

    def test_exact_margin_from_float_subtraction_is_not_flagged(self):
        from insurance.views import _rate_master_pi_po_rate_violations_qs

        flagged_ids = set(
            _rate_master_pi_po_rate_violations_qs("pi_od_rate", "po_od_rate").values_list("id", flat=True)
        )
        self.assertNotIn(self.exact_margin_row.id, flagged_ids)

    def test_a_genuinely_wrong_margin_is_still_flagged(self):
        from insurance.views import _rate_master_pi_po_rate_violations_qs

        flagged_ids = set(
            _rate_master_pi_po_rate_violations_qs("pi_od_rate", "po_od_rate").values_list("id", flat=True)
        )
        self.assertIn(self.genuine_violation_row.id, flagged_ids)


# Storage/cache overrides every Missing Make/Model test needs: MISFile carries
# a FileField, and the Add to Master write path calls cache.delete(), which
# would otherwise hit the (uncreated) DatabaseCache table.
MISSING_MM_OVERRIDES = dict(
    STORAGES={
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    },
    CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}},
)


def _make_model_reason(value):
    """The exact NO MATCH reason RULE 5a writes for an unresolvable make/model."""
    from insurance.mapping_engine import MAKE_MODEL_RULE_LABEL, make_model_unresolved_detail
    return f"Failed on: {MAKE_MODEL_RULE_LABEL} — {make_model_unresolved_detail(value)}"


class MakeModelFailureReasonConstantsTests(TestCase):
    """
    The Missing Make/Model page filters MISFailedRow.failure_reason on text the
    mapping engine wrote, sometimes months earlier. These guard the contract
    between the two: the message must stay byte-identical, the prefix must
    genuinely be a prefix, and the scope must stay narrowed to the one Rule 5a
    variant that adding to MakeModelMaster actually fixes.
    """

    def test_detail_text_is_byte_identical_to_the_historical_literal(self):
        # Hard-coded on purpose: this is what already-stored failure_reason rows
        # contain. If refactoring changes so much as a space, every historical
        # row silently drops off the page - so this literal must not be
        # regenerated from the constants it is checking.
        from insurance.mapping_engine import make_model_unresolved_detail
        self.assertEqual(
            make_model_unresolved_detail("yamaha alpha"),
            "Vehicle make/model 'yamaha alpha' did not share at least 2 words with any "
            "MakeModelMaster cluster entry — add it to MakeModelMaster or check the MIS value.",
        )

    def test_reason_prefix_matches_what_the_engine_builds(self):
        from insurance.mapping_engine import MAKE_MODEL_UNRESOLVED_REASON_PREFIX
        self.assertTrue(
            _make_model_reason("tata nexon ev").startswith(MAKE_MODEL_UNRESOLVED_REASON_PREFIX)
        )

    def test_value_pattern_round_trips_including_an_apostrophe(self):
        from insurance.mapping_engine import MAKE_MODEL_UNRESOLVED_VALUE_PATTERN
        for value in ["yamaha alpha", "bmw 3's series", "tata nexon ev", "maruti eeco(2012 - 2017)"]:
            match = MAKE_MODEL_UNRESOLVED_VALUE_PATTERN.match(_make_model_reason(value))
            self.assertIsNotNone(match, value)
            self.assertEqual(match.group(1), value)

    def test_the_generic_quoted_value_pattern_would_truncate_an_apostrophe(self):
        # Documents why MAKE_MODEL_UNRESOLVED_VALUE_PATTERN exists at all, so a
        # future "simplification" back to _QUOTED_VALUE_PATTERN fails loudly
        # instead of silently splitting one make/model into two page rows.
        from insurance.mapping_engine import _QUOTED_VALUE_PATTERN
        reason = _make_model_reason("bmw 3's series")
        self.assertEqual(_QUOTED_VALUE_PATTERN.search(reason).group(1), "bmw 3")

    def test_the_other_rule_5a_variants_are_out_of_scope(self):
        # Neither of these is fixed by adding to MakeModelMaster: the first is a
        # gap in the rate grid, the second is missing source data. Note the
        # first one DOES share the reason prefix - it opens with the same
        # "Vehicle make/model '<value>'" text and only diverges after the value.
        # That is why the page's filter needs the suffix marker as well, and why
        # this test asserts on the full pattern rather than the prefix alone.
        from insurance.mapping_engine import (
            MAKE_MODEL_RULE_LABEL, MAKE_MODEL_UNRESOLVED_REASON_PREFIX,
            MAKE_MODEL_UNRESOLVED_VALUE_PATTERN, MAKE_MODEL_UNRESOLVED_VALUE_SUFFIX,
        )
        resolved_but_uncovered = (
            f"Failed on: {MAKE_MODEL_RULE_LABEL} — Vehicle make/model 'yamaha alpha' resolved to "
            f"master group(s) [two_wheeler_all], but no remaining candidate rate row lists that "
            f"group in its vehicle-make cluster."
        )
        blank = (
            f"Failed on: {MAKE_MODEL_RULE_LABEL} — Vehicle make/model is blank on this policy, and "
            f"no remaining candidate rate row allows a blank vehicle-make cluster."
        )
        for reason in (resolved_but_uncovered, blank):
            self.assertIsNone(MAKE_MODEL_UNRESOLVED_VALUE_PATTERN.match(reason))
            self.assertFalse(
                reason.startswith(MAKE_MODEL_UNRESOLVED_REASON_PREFIX)
                and MAKE_MODEL_UNRESOLVED_VALUE_SUFFIX in reason
            )
        # The blank variant is excluded by the prefix alone; the sibling is not.
        self.assertFalse(blank.startswith(MAKE_MODEL_UNRESOLVED_REASON_PREFIX))
        self.assertTrue(resolved_but_uncovered.startswith(MAKE_MODEL_UNRESOLVED_REASON_PREFIX))

    def test_the_coverage_gap_parser_still_splits_the_refactored_reason(self):
        # _FAILED_ON_PATTERN stops at the FIRST em-dash; the detail sentence
        # contains a second one. build_coverage_gap_summary depends on that.
        from insurance.mapping_engine import _FAILED_ON_PATTERN, MAKE_MODEL_RULE_LABEL
        match = _FAILED_ON_PATTERN.match(_make_model_reason("yamaha alpha"))
        self.assertEqual(match.group(1).strip(), MAKE_MODEL_RULE_LABEL)
        self.assertTrue(match.group(2).startswith("Vehicle make/model 'yamaha alpha'"))


class MakeModelClusterIndexTests(TestCase):
    """
    build_make_model_cluster_index/resolve_make_model_with_index is a second
    implementation of Rule 5a's match decision, kept only because it hoists the
    master-side tokenization out of the per-term loop. It must never disagree
    with the engine's own fuzzy_match_make_model.
    """

    def test_index_agrees_with_fuzzy_match_make_model(self):
        from insurance.mapping_engine import (
            build_make_model_cluster_index, fuzzy_match_make_model, resolve_make_model_with_index,
        )
        from insurance.models import MakeModelMaster

        cases = [
            ("HONDA ACTIVA, TVS JUPITER", "yamaha alpha", False),
            ("HONDA ACTIVA, TVS JUPITER", "honda activa", True),
            ("HONDA ACTIVA6G", "honda activa 6g", True),      # digit-glue split
            ("HONDA CITY-1.5", "honda city 1.5", True),       # hyphen split
            ("HONDA ACTIVA, TVS JUPITER", "honda", False),    # <2 words can't match
            ("HONDA ACTIVA, TVS JUPITER", "honda jupiter", False),  # per-item, no cross-combining
            ("TATA NEXON EV PRIME", "tata nexon ev", True),
        ]
        for cluster, term, expected in cases:
            with self.subTest(cluster=cluster, term=term):
                MakeModelMaster.objects.all().delete()
                MakeModelMaster.objects.create(make_model_name="grp", make_model_cluster=cluster)
                index = build_make_model_cluster_index(
                    MakeModelMaster.objects.all(), "make_model_name", "make_model_cluster"
                )
                via_index = bool(resolve_make_model_with_index(index, term))
                via_engine = fuzzy_match_make_model(term, cluster)
                self.assertEqual(via_index, via_engine)
                self.assertEqual(via_index, expected)

    def test_rows_without_a_cluster_are_skipped(self):
        from insurance.mapping_engine import (
            build_make_model_cluster_index, resolve_make_model_with_index,
        )
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(make_model_name="empty", make_model_cluster="")
        MakeModelMaster.objects.create(make_model_name="null", make_model_cluster=None)
        MakeModelMaster.objects.create(make_model_name="commas", make_model_cluster=" , , ")
        index = build_make_model_cluster_index(
            MakeModelMaster.objects.all(), "make_model_name", "make_model_cluster"
        )
        # A master with nothing indexable must never resolve anything —
        # asserted through the public behaviour rather than the index's shape.
        self.assertEqual(resolve_make_model_with_index(index, "yamaha alpha"), set())


@override_settings(**MISSING_MM_OVERRIDES)
class MissingMakeModelAggregationTests(TestCase):
    """Missing Make/Model's grouping of MISFailedRow into one row per gap."""

    def setUp(self):
        from insurance.models import MISFile

        self.client = Client()
        Group.objects.get_or_create(name="Can_View_Missing_Make_Model")
        self.user = User.objects.create_user(username="ops", password="a-strong-test-password-1")
        self.user.groups.add(Group.objects.get(name="Can_View_Missing_Make_Model"))
        self.client.force_login(self.user)

        self.mis_file = MISFile.objects.create(status="COMPLETED", uploaded_file="mis/jan.xlsx")
        self._row_id = 0

    def _fail(self, make, model, product="Two Wheeler", sub_product="Scooter",
              insurer="Acme General", vehicle_class="", reason=None, status_key="NO_MATCH",
              payload=None, mis_file=None):
        """One MISFailedRow shaped the way mapping_engine writes them."""
        from insurance.models import MISFailedRow
        self._row_id += 1
        value = f"{make} {model}".strip().lower()
        if payload is None:
            payload = {
                "Policy: vehicle make": make,
                "Policy: model": model,
                "Policy: vehproduct": product,
                "Policy: sub product": sub_product,
                "Policy: insurance company": insurer,
            }
            if vehicle_class:
                payload["Policy: vehicle class"] = vehicle_class
        return MISFailedRow.objects.create(
            mis_file=mis_file or self.mis_file,
            row_id=self._row_id,
            status_key=status_key,
            mapping_status="❌ NO MATCH",
            failure_reason=reason if reason is not None else _make_model_reason(value),
            insurer=insurer,
            payload=payload,
        )

    def _groups(self):
        from insurance.views import _missing_make_model_groups
        return _missing_make_model_groups()

    def test_casing_variants_collapse_into_one_row(self):
        self._fail("YAMAHA", "ALPHA")
        self._fail("yamaha", "alpha")
        groups = self._groups()
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["value"], "yamaha alpha")
        self.assertEqual(groups[0]["policy_count"], 2)
        # First non-blank, oldest row first - so display casing is stable.
        self.assertEqual(groups[0]["make"], "YAMAHA")

    def test_payload_keys_are_read_case_insensitively(self):
        # payload keys are the uploaded file's own header text, so casing varies
        # between files and a DB-level JSON lookup would miss these.
        self._fail("TATA", "NEXON EV", payload={
            "Policy: Vehicle Make": "TATA",
            "policy: model": "NEXON EV",
            "POLICY: VEHPRODUCT": "Private Car",
            "  Policy: Sub Product  ": "Hatchback",
        })
        groups = self._groups()
        self.assertEqual(groups[0]["make"], "TATA")
        self.assertEqual(groups[0]["model"], "NEXON EV")
        self.assertEqual(groups[0]["product"], "Private Car")
        self.assertEqual(groups[0]["sub_product"], "Hatchback")

    def test_distinct_product_splits_the_group(self):
        self._fail("YAMAHA", "ALPHA", product="Two Wheeler")
        self._fail("YAMAHA", "ALPHA", product="GCV")
        self.assertEqual(len(self._groups()), 2)

    def test_distinct_insurer_splits_the_group(self):
        self._fail("YAMAHA", "ALPHA", insurer="Acme General")
        self._fail("YAMAHA", "ALPHA", insurer="Zenith Insurance")
        self.assertEqual(len(self._groups()), 2)

    def test_distinct_sub_product_splits_the_group(self):
        # Pins the deliberate choice to key on sub product as well as product,
        # because the right cluster to add a make/model to depends on both. If
        # that ever turns out to fragment the page too much, this is the single
        # test that says so.
        self._fail("YAMAHA", "ALPHA", product="GCV", sub_product="3W")
        self._fail("YAMAHA", "ALPHA", product="GCV", sub_product="4W")
        self.assertEqual(len(self._groups()), 2)

    def test_distinct_vehicle_class_splits_the_group(self):
        # Same reasoning as sub product: RULE 2b (vehicle class) narrows
        # current_grid before RULE 5a runs, so whether a master group actually
        # fixes this failure can depend on vehicle class too.
        self._fail("YAMAHA", "ALPHA", vehicle_class="Bike")
        self._fail("YAMAHA", "ALPHA", vehicle_class="Scooter")
        self.assertEqual(len(self._groups()), 2)

    def test_blank_model_payload_key_is_tolerated(self):
        # _extract_failed_rows_from_df skips NaN cells entirely, so a policy
        # with no model has no 'Policy: model' key at all - while the engine
        # still saw (and reported) the make.
        self._fail("YAMAHA", "", payload={"Policy: vehicle make": "YAMAHA"})
        groups = self._groups()
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["value"], "yamaha")
        self.assertEqual(groups[0]["model"], "")

    def test_a_row_with_no_make_or_model_payload_still_shows_the_failing_value(self):
        self._fail("YAMAHA", "ALPHA", payload={"Policy: vehproduct": "Two Wheeler"})
        self.assertEqual(self._groups()[0]["make"], "yamaha alpha")

    def test_other_failure_reasons_are_excluded(self):
        from insurance.mapping_engine import MAKE_MODEL_RULE_LABEL
        self._fail("YAMAHA", "ALPHA")                      # in scope
        self._fail("TATA", "NEXON", reason=(               # Rule 5b, out of scope
            "Failed on: Policy: rto no — RTO 'HR51' resolved to master group(s) [allindia], but "
            "no remaining candidate rate row lists that group in its RTO cluster."
        ))
        self._fail("HONDA", "CITY", reason=(               # sibling 5a, out of scope
            f"Failed on: {MAKE_MODEL_RULE_LABEL} — Vehicle make/model 'honda city' resolved to "
            f"master group(s) [private_car_all], but no remaining candidate rate row lists that "
            f"group in its vehicle-make cluster."
        ))
        groups = self._groups()
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["value"], "yamaha alpha")

    def test_multiple_matches_rows_are_excluded(self):
        self._fail("YAMAHA", "ALPHA", status_key="MULTIPLE_MATCHES")
        self.assertEqual(self._groups(), [])

    def test_policy_count_and_file_count(self):
        from insurance.models import MISFile
        other_file = MISFile.objects.create(status="COMPLETED", uploaded_file="mis/feb.xlsx")
        self._fail("YAMAHA", "ALPHA")
        self._fail("YAMAHA", "ALPHA")
        self._fail("YAMAHA", "ALPHA", mis_file=other_file)
        group = self._groups()[0]
        self.assertEqual(group["policy_count"], 3)
        self.assertEqual(group["file_count"], 2)

    def test_rows_are_sorted_by_policy_count(self):
        self._fail("TATA", "NEXON EV")
        self._fail("YAMAHA", "ALPHA")
        self._fail("YAMAHA", "ALPHA")
        self.assertEqual([g["value"] for g in self._groups()], ["yamaha alpha", "tata nexon ev"])


@override_settings(**MISSING_MM_OVERRIDES)
class MissingMakeModelPageTests(MissingMakeModelAggregationTests):
    """The page itself: resolution status, filters, export."""

    def _wire_active_rate(self, insurer, product, sub_product, make_name, vehicle_class=None):
        """
        An ACTIVE, non-deleted Rate Master row that lists `make_name` in its
        new_vehicle_makes cluster for (insurer, product, sub_product) -- Step 2
        of RULE 5a's chain, and what _active_vehicle_make_scope reads.
        vehicle_class=None leaves make_model_class unset, i.e. RULE 2b's
        NA-wildcard row (matches any MIS vehicle class).
        """
        from insurance.models import MakeModelClassMaster, ProductMaster, RateMaster, SubProductMaster
        product_obj, _ = ProductMaster.objects.get_or_create(name=product)
        sub_product_obj, _ = SubProductMaster.objects.get_or_create(name=sub_product)
        class_obj = None
        if vehicle_class:
            class_obj, _ = MakeModelClassMaster.objects.get_or_create(name=vehicle_class)
        return RateMaster.objects.create(
            insurance_company=insurer, product=product_obj, sub_product=sub_product_obj,
            make_model_class=class_obj, new_vehicle_makes=make_name,
            status="ACTIVE", is_deleted="NO",
        )

    def test_a_value_already_in_a_cluster_is_marked_resolved(self):
        from insurance.models import MakeModelMaster
        self._fail("YAMAHA", "ALPHA")
        MakeModelMaster.objects.create(
            make_model_name="two_wheeler_all", make_model_cluster="HONDA ACTIVA, YAMAHA ALPHA"
        )
        # Resolution requires an ACTIVE rate row for the SAME insurer/product/
        # sub product to actually reference the matched master group (Step 2) --
        # a bare Step-1 word-overlap match is not enough. See
        # test_a_cluster_match_unused_by_this_insurer_is_not_resolved.
        self._wire_active_rate("Acme General", "Two Wheeler", "Scooter", "two_wheeler_all")
        response = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(response.context["page_obj"])[0]
        self.assertTrue(group["resolved"])
        self.assertEqual(group["resolved_names"], ["two_wheeler_all"])

    def test_a_cluster_match_unused_by_this_insurer_is_not_resolved(self):
        # Reported live-site bug: a Liberty / Private Car / SAOD failure fuzzy-
        # matched (Step 1, word-overlap only) a MakeModelMaster group that
        # Liberty's own grids never reference in new_vehicle_makes -- only some
        # OTHER insurer's grid used it. That must not count as Resolved, since
        # reprocessing wouldn't actually map the policy; it would just trade
        # this failure for "resolved to master group(s) [...] but no candidate
        # rate row".
        self._fail("HYUNDAI", "ER/SX 1.2 CNG MT-SUV", product="Private Car",
                    sub_product="SAOD", insurer="Liberty General")
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(
            make_model_name="royal_sep26_hyundai_group",
            make_model_cluster="HYUNDAI ER/SX 1.2 CNG MT-SUV",
        )
        # Wired to a different insurer entirely -- Liberty's own grid never
        # lists this master group.
        self._wire_active_rate("Royal Sundaram", "Private Car", "SAOD", "royal_sep26_hyundai_group")

        response = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(response.context["page_obj"])[0]
        self.assertFalse(group["resolved"])
        self.assertEqual(group["resolved_names"], [])

    def test_a_cluster_match_unused_by_this_product_is_not_resolved(self):
        # Same scoping bug, but split on product/sub product instead of
        # insurer: the master group is wired for this insurer under a
        # different product, so it still doesn't fix THIS group.
        self._fail("HYUNDAI", "ER/SX 1.2 CNG MT-SUV", product="Private Car",
                    sub_product="SAOD", insurer="Liberty General")
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(
            make_model_name="hyundai_group", make_model_cluster="HYUNDAI ER/SX 1.2 CNG MT-SUV"
        )
        self._wire_active_rate("Liberty General", "Private Car", "Comprehensive", "hyundai_group")

        response = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(response.context["page_obj"])[0]
        self.assertFalse(group["resolved"])

    def _mark(self, value, product="Two Wheeler", sub_product="Scooter", insurer="Acme General",
              vehicle_class="", **overrides):
        data = {
            "mm_mark_value": value, "mm_mark_product": product, "mm_mark_sub_product": sub_product,
            "mm_mark_insurer": insurer, "mm_mark_vehicle_class": vehicle_class,
        }
        data.update(overrides)
        return self.client.post(reverse("mark_missing_make_model_resolved"), data)

    def _unmark(self, value, product="Two Wheeler", sub_product="Scooter", insurer="Acme General",
                vehicle_class="", **overrides):
        data = {
            "mm_mark_value": value, "mm_mark_product": product, "mm_mark_sub_product": sub_product,
            "mm_mark_insurer": insurer, "mm_mark_vehicle_class": vehicle_class,
        }
        data.update(overrides)
        return self.client.post(reverse("unmark_missing_make_model_resolved"), data)

    def test_marking_a_row_resolved_manually_needs_no_cluster_write(self):
        from insurance.models import MakeModelMaster, MissingMakeModelManualResolution
        self._fail("YAMAHA", "ALPHA")
        response = self._mark("yamaha alpha")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MakeModelMaster.objects.count(), 0)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 1)

        page = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(page.context["page_obj"])[0]
        self.assertTrue(group["resolved"])
        self.assertTrue(group["resolved_manual"])
        self.assertFalse(group["resolved_live"])

    def test_marking_twice_is_idempotent(self):
        from insurance.models import AuditLog, MissingMakeModelManualResolution
        self._fail("YAMAHA", "ALPHA")
        self._mark("yamaha alpha")
        self._mark("yamaha alpha")
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 1)
        self.assertEqual(AuditLog.objects.filter(action="MAKE MODEL MANUALLY RESOLVED").count(), 1)

    def test_unmarking_reverts_a_purely_manual_resolution(self):
        from insurance.models import MissingMakeModelManualResolution
        self._fail("YAMAHA", "ALPHA")
        self._mark("yamaha alpha")
        response = self._unmark("yamaha alpha")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 0)

        page = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(page.context["page_obj"])[0]
        self.assertFalse(group["resolved"])

    def test_unmarking_does_not_undo_a_live_match(self):
        # A row can be both live-resolved (real cluster match) and separately
        # marked manually (e.g. by mistake, or before the live match existed).
        # Undoing the manual mark must not hide a genuinely live-resolved row.
        from insurance.models import MakeModelMaster, MissingMakeModelManualResolution
        self._fail("YAMAHA", "ALPHA")
        MakeModelMaster.objects.create(make_model_name="two_wheeler_all", make_model_cluster="YAMAHA ALPHA")
        self._wire_active_rate("Acme General", "Two Wheeler", "Scooter", "two_wheeler_all")
        self._mark("yamaha alpha")
        self._unmark("yamaha alpha")
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 0)

        page = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(page.context["page_obj"])[0]
        self.assertTrue(group["resolved"])
        self.assertTrue(group["resolved_live"])

    def test_unmarking_a_row_that_was_never_marked_is_a_no_op(self):
        from insurance.models import MissingMakeModelManualResolution
        self._fail("YAMAHA", "ALPHA")
        response = self._unmark("yamaha alpha")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 0)

    def test_marking_a_get_redirects_without_writing(self):
        from insurance.models import MissingMakeModelManualResolution
        self._fail("YAMAHA", "ALPHA")
        response = self.client.get(reverse("mark_missing_make_model_resolved"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 0)

    def test_a_cluster_match_for_a_different_vehicle_class_is_not_resolved(self):
        # Same scoping bug again, this time on RULE 2b's own dimension: the
        # master group is wired for the right insurer/product/sub product, but
        # under a DIFFERENT, specific vehicle class (not NA-wildcard) -- so it
        # still doesn't fix a "Bike" failure.
        self._fail("YAMAHA", "ALPHA", vehicle_class="Bike")
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(
            make_model_name="two_wheeler_all", make_model_cluster="YAMAHA ALPHA"
        )
        self._wire_active_rate("Acme General", "Two Wheeler", "Scooter", "two_wheeler_all",
                                vehicle_class="Scooter")

        response = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(response.context["page_obj"])[0]
        self.assertFalse(group["resolved"])

    def test_an_exact_vehicle_class_match_is_resolved(self):
        self._fail("YAMAHA", "ALPHA", vehicle_class="Bike")
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(
            make_model_name="two_wheeler_all", make_model_cluster="YAMAHA ALPHA"
        )
        self._wire_active_rate("Acme General", "Two Wheeler", "Scooter", "two_wheeler_all",
                                vehicle_class="Bike")

        response = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(response.context["page_obj"])[0]
        self.assertTrue(group["resolved"])

    def test_a_na_wildcard_rate_row_resolves_any_vehicle_class(self):
        # match_vehicle_class treats a blank/NA make_model_class as a wildcard
        # that passes regardless of the MIS row's class -- _active_vehicle_
        # make_scope must honour that too, not just an exact class match.
        self._fail("YAMAHA", "ALPHA", vehicle_class="Bike")
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(
            make_model_name="two_wheeler_all", make_model_cluster="YAMAHA ALPHA"
        )
        self._wire_active_rate("Acme General", "Two Wheeler", "Scooter", "two_wheeler_all")

        response = self.client.get(reverse("missing_make_model"), {"status": "all"})
        group = list(response.context["page_obj"])[0]
        self.assertTrue(group["resolved"])

    def test_default_status_filter_hides_resolved_rows(self):
        from insurance.models import MakeModelMaster
        self._fail("YAMAHA", "ALPHA")
        self._fail("TATA", "NEXON EV")
        MakeModelMaster.objects.create(
            make_model_name="two_wheeler_all", make_model_cluster="YAMAHA ALPHA"
        )
        self._wire_active_rate("Acme General", "Two Wheeler", "Scooter", "two_wheeler_all")
        response = self.client.get(reverse("missing_make_model"))
        self.assertEqual(response.context["selected"]["status"], "missing")
        self.assertEqual([g["value"] for g in response.context["page_obj"]], ["tata nexon ev"])
        self.assertEqual(response.context["total_missing"], 1)
        self.assertEqual(response.context["total_resolved"], 1)

        resolved_only = self.client.get(reverse("missing_make_model"), {"status": "resolved"})
        self.assertEqual([g["value"] for g in resolved_only.context["page_obj"]], ["yamaha alpha"])

        everything = self.client.get(reverse("missing_make_model"), {"status": "all"})
        self.assertEqual(len(list(everything.context["page_obj"])), 2)

    def test_insurer_product_and_search_filters(self):
        self._fail("YAMAHA", "ALPHA", product="Two Wheeler", insurer="Acme General")
        self._fail("TATA", "NEXON EV", product="Private Car", insurer="Zenith Insurance")

        by_insurer = self.client.get(reverse("missing_make_model"), {"insurer": "Zenith Insurance"})
        self.assertEqual([g["value"] for g in by_insurer.context["page_obj"]], ["tata nexon ev"])

        by_product = self.client.get(reverse("missing_make_model"), {"product": "Two Wheeler"})
        self.assertEqual([g["value"] for g in by_product.context["page_obj"]], ["yamaha alpha"])

        by_search = self.client.get(reverse("missing_make_model"), {"q": "NEXON"})
        self.assertEqual([g["value"] for g in by_search.context["page_obj"]], ["tata nexon ev"])

    def test_filter_dropdowns_are_built_from_the_unfiltered_aggregate(self):
        # Applying a filter must not remove its own option from the list.
        self._fail("YAMAHA", "ALPHA", product="Two Wheeler", insurer="Acme General")
        self._fail("TATA", "NEXON EV", product="Private Car", insurer="Zenith Insurance")
        response = self.client.get(reverse("missing_make_model"), {"insurer": "Zenith Insurance"})
        self.assertEqual(response.context["insurer_list"], ["Acme General", "Zenith Insurance"])
        self.assertEqual(response.context["product_list"], ["Private Car", "Two Wheeler"])

    def test_an_invalid_date_is_ignored_rather_than_erroring(self):
        self._fail("YAMAHA", "ALPHA")
        response = self.client.get(reverse("missing_make_model"), {"date_from": "not-a-date"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_all"], 1)
        self.assertEqual(response.context["selected"]["date_from"], "")

    def test_export_honours_filters_and_ignores_pagination(self):
        self._fail("YAMAHA", "ALPHA", insurer="Acme General")
        self._fail("TATA", "NEXON EV", insurer="Zenith Insurance")
        response = self.client.get(
            reverse("export_missing_make_model_xlsx"),
            {"insurer": "Zenith Insurance", "page": "99"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("missing_make_model.xlsx", response["Content-Disposition"])

        import io
        from openpyxl import load_workbook
        rows = list(load_workbook(io.BytesIO(response.content)).active.values)
        self.assertEqual(len(rows), 2)                     # header + the one match
        self.assertEqual(rows[1][3], "tata nexon ev")

    def test_a_user_without_the_group_is_refused(self):
        other = User.objects.create_user(username="nobody", password="a-strong-test-password-2")
        client = Client()
        client.force_login(other)
        self.assertEqual(client.get(reverse("missing_make_model")).status_code, 403)


@override_settings(**MISSING_MM_OVERRIDES)
class AddMissingMakeModelToMasterTests(TestCase):
    """The one write path into MakeModelMaster.make_model_cluster outside the importer."""

    def setUp(self):
        from insurance.models import MakeModelMaster

        self.client = Client()
        Group.objects.get_or_create(name="Can_View_Missing_Make_Model")
        self.user = User.objects.create_user(username="ops", password="a-strong-test-password-1")
        self.user.groups.add(Group.objects.get(name="Can_View_Missing_Make_Model"))
        self.client.force_login(self.user)

        self.master = MakeModelMaster.objects.create(
            make_model_name="two_wheeler_all", make_model_cluster="HONDA ACTIVA, TVS JUPITER"
        )
        self.url = reverse("add_missing_make_model_to_master")

    def _post(self, **overrides):
        data = {
            "make_model_value": "yamaha alpha",
            "target": "existing",
            "master_id": str(self.master.id),
        }
        data.update(overrides)
        return self.client.post(self.url, data)

    def test_appends_to_an_existing_cluster_preserving_existing_text(self):
        response = self._post()
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER, YAMAHA ALPHA")
        self.assertEqual(response.status_code, 302)

    def test_the_value_is_stored_upper_cased(self):
        # fuzzy_match_make_model upper-cases both sides, so casing can't change
        # matching - it is stored upper only to read like its neighbours.
        self._post(make_model_value="tata nexon ev")
        self.master.refresh_from_db()
        self.assertTrue(self.master.make_model_cluster.endswith("TATA NEXON EV"))

    def test_the_appended_value_actually_resolves_afterwards(self):
        from insurance.mapping_engine import fuzzy_match_make_model
        self._post()
        self.master.refresh_from_db()
        self.assertTrue(fuzzy_match_make_model("yamaha alpha", self.master.make_model_cluster))

    def test_a_duplicate_item_is_a_no_op(self):
        from insurance.models import AuditLog
        self.master.make_model_cluster = "HONDA ACTIVA, YAMAHA ALPHA"
        self.master.save()
        self._post(force="1")                       # force past the already-resolves guard
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, YAMAHA ALPHA")
        self.assertFalse(AuditLog.objects.filter(action="MAKE MODEL CLUSTER ADD").exists())

    def test_a_trailing_comma_does_not_produce_a_double_separator(self):
        self.master.make_model_cluster = "HONDA ACTIVA, "
        self.master.save()
        self._post()
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, YAMAHA ALPHA")

    def test_creates_a_new_master_row(self):
        from insurance.models import MakeModelMaster
        self._post(target="new", new_name="private_car_ev", make_model_value="tata nexon ev")
        created = MakeModelMaster.objects.get(make_model_name="private_car_ev")
        self.assertEqual(created.make_model_cluster, "TATA NEXON EV")

    def test_a_duplicate_new_name_is_rejected_case_insensitively(self):
        from insurance.models import MakeModelMaster
        self._post(target="new", new_name="TWO_WHEELER_ALL")
        self.assertEqual(MakeModelMaster.objects.count(), 1)

    def test_a_blank_new_name_is_rejected(self):
        from insurance.models import MakeModelMaster
        self._post(target="new", new_name="   ")
        self.assertEqual(MakeModelMaster.objects.count(), 1)

    def test_a_comma_in_the_value_is_replaced_with_a_space(self):
        # A comma would split the append into two cluster items, and a 1-word
        # item can never reach the 2-shared-word threshold.
        self._post(make_model_value="yamaha, alpha")
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER, YAMAHA ALPHA")

    def test_an_empty_value_is_rejected(self):
        self._post(make_model_value="   ")
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER")

    def test_an_overlong_value_is_rejected(self):
        self._post(make_model_value="a b " * 200)
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER")

    def test_a_missing_master_row_is_rejected(self):
        response = self._post(master_id="999999")
        self.assertEqual(response.status_code, 302)
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER")

    def test_a_value_can_be_added_to_multiple_existing_masters_at_once(self):
        from insurance.models import MakeModelMaster
        other = MakeModelMaster.objects.create(make_model_name="private_car_all", make_model_cluster="TATA NEXON")
        self._post(master_id=[str(self.master.id), str(other.id)])
        self.master.refresh_from_db()
        other.refresh_from_db()
        self.assertIn("YAMAHA ALPHA", self.master.make_model_cluster)
        self.assertIn("YAMAHA ALPHA", other.make_model_cluster)

    def test_multi_master_add_writes_one_audit_log_entry_naming_both(self):
        from insurance.models import AuditLog, MakeModelMaster
        other = MakeModelMaster.objects.create(make_model_name="private_car_all", make_model_cluster="TATA NEXON")
        self._post(master_id=[str(self.master.id), str(other.id)])
        entry = AuditLog.objects.get(action="MAKE MODEL CLUSTER ADD")
        self.assertIn("two_wheeler_all", entry.details)
        self.assertIn("private_car_all", entry.details)

    def test_multi_master_add_skips_a_row_where_it_is_already_present(self):
        from insurance.models import MakeModelMaster
        self.master.make_model_cluster = "HONDA ACTIVA, YAMAHA ALPHA"
        self.master.save()
        other = MakeModelMaster.objects.create(make_model_name="private_car_all", make_model_cluster="TATA NEXON")
        self._post(force="1", master_id=[str(self.master.id), str(other.id)])
        self.master.refresh_from_db()
        other.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, YAMAHA ALPHA")
        self.assertIn("YAMAHA ALPHA", other.make_model_cluster)

    def test_a_value_already_in_every_selected_master_is_a_no_op(self):
        from insurance.models import AuditLog, MakeModelMaster
        self.master.make_model_cluster = "HONDA ACTIVA, YAMAHA ALPHA"
        self.master.save()
        other = MakeModelMaster.objects.create(make_model_name="private_car_all", make_model_cluster="YAMAHA ALPHA")
        self._post(force="1", master_id=[str(self.master.id), str(other.id)])
        self.assertFalse(AuditLog.objects.filter(action="MAKE MODEL CLUSTER ADD").exists())

    def test_an_id_that_matches_nothing_among_several_is_silently_dropped(self):
        self._post(master_id=[str(self.master.id), "999999"])
        self.master.refresh_from_db()
        self.assertIn("YAMAHA ALPHA", self.master.make_model_cluster)

    def test_an_unknown_target_is_rejected(self):
        self._post(target="")
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER")

    def test_an_already_resolving_value_is_refused_without_force(self):
        # Adding it to a SECOND cluster would make these policies resolve to two
        # master groups - MULTIPLE MATCHES instead of a fix.
        from insurance.models import MakeModelMaster
        other = MakeModelMaster.objects.create(
            make_model_name="scooters", make_model_cluster="YAMAHA ALPHA"
        )
        self._post()
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER")
        other.refresh_from_db()
        self.assertEqual(other.make_model_cluster, "YAMAHA ALPHA")

    def test_an_already_resolving_value_is_accepted_with_force(self):
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(make_model_name="scooters", make_model_cluster="YAMAHA ALPHA")
        self._post(force="1")
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER, YAMAHA ALPHA")

    def test_an_audit_log_entry_is_written(self):
        from insurance.models import AuditLog
        self._post(policy_count="412")
        entry = AuditLog.objects.get(action="MAKE MODEL CLUSTER ADD")
        self.assertEqual(entry.user, self.user)
        self.assertIn("YAMAHA ALPHA", entry.details)
        self.assertIn("two_wheeler_all", entry.details)
        self.assertIn("412", entry.details)

    def test_the_rate_form_choices_cache_is_invalidated(self):
        from django.core.cache import cache
        from insurance.views import RTO_MAKE_CHOICES_CACHE_KEY
        cache.set(RTO_MAKE_CHOICES_CACHE_KEY, {"rtos": [], "makes": []}, 600)
        self._post()
        self.assertIsNone(cache.get(RTO_MAKE_CHOICES_CACHE_KEY))

    def test_a_get_redirects_without_writing(self):
        response = self.client.get(self.url, {"make_model_value": "yamaha alpha"})
        self.assertEqual(response.status_code, 302)
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER")

    def test_the_redirect_preserves_the_page_filters(self):
        response = self._post(insurer="Acme General", product="Two Wheeler", status="all", q="alpha")
        self.assertEqual(response.status_code, 302)
        self.assertIn("insurer=Acme+General", response["Location"])
        self.assertIn("product=Two+Wheeler", response["Location"])
        self.assertIn("status=all", response["Location"])
        self.assertIn("q=alpha", response["Location"])

    def test_a_user_without_the_group_is_refused(self):
        other = User.objects.create_user(username="nobody", password="a-strong-test-password-2")
        client = Client()
        client.force_login(other)
        response = client.post(self.url, {
            "make_model_value": "yamaha alpha", "target": "existing", "master_id": str(self.master.id),
        })
        self.assertEqual(response.status_code, 403)
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA, TVS JUPITER")


@override_settings(**MISSING_MM_OVERRIDES)
class BulkMissingMakeModelActionsTests(TestCase):
    """
    Bulk select-and-assign ("Add Selected to Master") and bulk "Mark Selected
    Resolved" -- the checkbox-driven versions of the single-row actions above,
    for when several distinct Missing Make/Model rows should all get the same
    treatment in one request.
    """

    def setUp(self):
        from insurance.models import MakeModelMaster

        self.client = Client()
        Group.objects.get_or_create(name="Can_View_Missing_Make_Model")
        self.user = User.objects.create_user(username="ops", password="a-strong-test-password-1")
        self.user.groups.add(Group.objects.get(name="Can_View_Missing_Make_Model"))
        self.client.force_login(self.user)

        self.master = MakeModelMaster.objects.create(
            make_model_name="two_wheeler_all", make_model_cluster="HONDA ACTIVA"
        )
        self.add_url = reverse("bulk_add_missing_make_model_to_master")
        self.mark_url = reverse("bulk_mark_missing_make_model_resolved")

    def _key(self, value, product="Two Wheeler", sub_product="Scooter", insurer="Acme General",
             vehicle_class=""):
        return json.dumps({
            "value": value, "product": product, "sub_product": sub_product,
            "insurer": insurer, "vehicle_class": vehicle_class,
        })

    def test_bulk_add_appends_every_selected_value_to_one_master(self):
        response = self.client.post(self.add_url, {
            "group_keys": [self._key("yamaha alpha"), self._key("tvs jupiter zx")],
            "target": "existing", "master_id": str(self.master.id),
        })
        self.assertEqual(response.status_code, 302)
        self.master.refresh_from_db()
        self.assertIn("YAMAHA ALPHA", self.master.make_model_cluster)
        self.assertIn("TVS JUPITER ZX", self.master.make_model_cluster)

    def test_bulk_add_dedupes_a_repeated_value(self):
        self.client.post(self.add_url, {
            "group_keys": [self._key("yamaha alpha"), self._key("yamaha alpha")],
            "target": "existing", "master_id": str(self.master.id),
        })
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster.count("YAMAHA ALPHA"), 1)

    def test_bulk_add_creates_a_new_master_row(self):
        from insurance.models import MakeModelMaster
        self.client.post(self.add_url, {
            "group_keys": [self._key("tata nexon ev")],
            "target": "new", "new_name": "private_car_ev",
        })
        created = MakeModelMaster.objects.get(make_model_name="private_car_ev")
        self.assertEqual(created.make_model_cluster, "TATA NEXON EV")

    def test_bulk_add_to_multiple_existing_masters_at_once(self):
        from insurance.models import MakeModelMaster
        other = MakeModelMaster.objects.create(make_model_name="private_car_all", make_model_cluster="TATA NEXON")
        self.client.post(self.add_url, {
            "group_keys": [self._key("yamaha alpha")],
            "target": "existing", "master_id": [str(self.master.id), str(other.id)],
        })
        self.master.refresh_from_db()
        other.refresh_from_db()
        self.assertIn("YAMAHA ALPHA", self.master.make_model_cluster)
        self.assertIn("YAMAHA ALPHA", other.make_model_cluster)

    def test_bulk_add_skips_a_value_that_already_resolves_elsewhere_without_force(self):
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(make_model_name="scooters", make_model_cluster="YAMAHA ALPHA")
        self.client.post(self.add_url, {
            "group_keys": [self._key("yamaha alpha"), self._key("tvs jupiter zx")],
            "target": "existing", "master_id": str(self.master.id),
        })
        self.master.refresh_from_db()
        self.assertNotIn("YAMAHA ALPHA", self.master.make_model_cluster)
        self.assertIn("TVS JUPITER ZX", self.master.make_model_cluster)

    def test_bulk_add_includes_a_skipped_value_with_force(self):
        from insurance.models import MakeModelMaster
        MakeModelMaster.objects.create(make_model_name="scooters", make_model_cluster="YAMAHA ALPHA")
        self.client.post(self.add_url, {
            "group_keys": [self._key("yamaha alpha")],
            "target": "existing", "master_id": str(self.master.id), "force": "1",
        })
        self.master.refresh_from_db()
        self.assertIn("YAMAHA ALPHA", self.master.make_model_cluster)

    def test_bulk_add_with_no_selection_is_rejected(self):
        response = self.client.post(self.add_url, {"target": "existing", "master_id": str(self.master.id)})
        self.assertEqual(response.status_code, 302)
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA")

    def test_bulk_add_a_get_redirects_without_writing(self):
        response = self.client.get(self.add_url)
        self.assertEqual(response.status_code, 302)
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA")

    def test_bulk_add_writes_one_audit_log_entry(self):
        from insurance.models import AuditLog
        self.client.post(self.add_url, {
            "group_keys": [self._key("yamaha alpha"), self._key("tvs jupiter zx")],
            "target": "existing", "master_id": str(self.master.id),
        })
        entries = AuditLog.objects.filter(action="MAKE MODEL CLUSTER BULK ADD")
        self.assertEqual(entries.count(), 1)
        self.assertIn("2", entries.first().details)

    def test_bulk_add_invalidates_the_rate_form_choices_cache(self):
        from django.core.cache import cache
        from insurance.views import RTO_MAKE_CHOICES_CACHE_KEY
        cache.set(RTO_MAKE_CHOICES_CACHE_KEY, {"rtos": [], "makes": []}, 600)
        self.client.post(self.add_url, {
            "group_keys": [self._key("yamaha alpha")],
            "target": "existing", "master_id": str(self.master.id),
        })
        self.assertIsNone(cache.get(RTO_MAKE_CHOICES_CACHE_KEY))

    def test_bulk_mark_resolves_every_selected_row(self):
        from insurance.models import MissingMakeModelManualResolution
        response = self.client.post(self.mark_url, {
            "group_keys": [self._key("yamaha alpha"), self._key("tata nexon ev", product="Private Car")],
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 2)

    def test_bulk_mark_is_idempotent_for_an_already_marked_row(self):
        from insurance.models import MissingMakeModelManualResolution
        self.client.post(self.mark_url, {"group_keys": [self._key("yamaha alpha")]})
        self.client.post(self.mark_url, {"group_keys": [self._key("yamaha alpha")]})
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 1)

    def test_bulk_mark_with_no_selection_is_rejected(self):
        from insurance.models import MissingMakeModelManualResolution
        response = self.client.post(self.mark_url, {})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 0)

    def test_bulk_mark_a_get_redirects_without_writing(self):
        from insurance.models import MissingMakeModelManualResolution
        response = self.client.get(self.mark_url)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 0)

    def test_bulk_mark_writes_one_audit_log_entry(self):
        from insurance.models import AuditLog
        self.client.post(self.mark_url, {
            "group_keys": [self._key("yamaha alpha"), self._key("tata nexon ev", product="Private Car")],
        })
        entries = AuditLog.objects.filter(action="MAKE MODEL MANUALLY RESOLVED")
        self.assertEqual(entries.count(), 1)

    def test_a_user_without_the_group_is_refused_on_every_new_endpoint(self):
        from insurance.models import MissingMakeModelManualResolution
        other = User.objects.create_user(username="nobody", password="a-strong-test-password-2")
        client = Client()
        client.force_login(other)
        urls = [
            self.add_url, self.mark_url,
            reverse("mark_missing_make_model_resolved"), reverse("unmark_missing_make_model_resolved"),
        ]
        for url in urls:
            response = client.post(url, {"group_keys": [self._key("yamaha alpha")]})
            self.assertEqual(response.status_code, 403, url)
        self.assertEqual(MissingMakeModelManualResolution.objects.count(), 0)
        self.master.refresh_from_db()
        self.assertEqual(self.master.make_model_cluster, "HONDA ACTIVA")
